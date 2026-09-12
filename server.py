#!/usr/bin/env python3
"""
server.py - Bridge server for Plagiarism Detection & Document Similarity Analysis.
Handles HTTP REST API, multi-format text extraction (.txt, .pdf, .docx, .pptx),
and delegates core algorithmic analysis to the authoritative C11 DSA engine.
Standard C engine receives normalized UTF-8 plain text.
"""

import http.server
import socketserver
import subprocess
import json
import os
import sys
import mimetypes
import base64
import uuid
import hashlib
from urllib.parse import urlparse

# Document parsing libraries
import pymupdf
import docx
import pptx

HOST = "0.0.0.0"
DEFAULT_PORT = 8080
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UI_DIR = os.path.join(BASE_DIR, "ui")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
ORIGINALS_DIR = os.path.join(UPLOADS_DIR, "originals")
EXTRACTED_DIR = os.path.join(UPLOADS_DIR, "extracted")
REPORTS_DIR = os.path.join(BASE_DIR, "reports")
DOCS_DIR = os.path.join(BASE_DIR, "documents")

# Locate C executable relative to project root
EXECUTABLE_NAME = "plagiarism_detector.exe" if sys.platform.startswith("win") else "plagiarism_detector"
C_EXECUTABLE = os.path.join(BASE_DIR, EXECUTABLE_NAME)

# Ensure writable directories exist
os.makedirs(ORIGINALS_DIR, exist_ok=True)
os.makedirs(EXTRACTED_DIR, exist_ok=True)
os.makedirs(REPORTS_DIR, exist_ok=True)
os.makedirs(DOCS_DIR, exist_ok=True)

MAX_UPLOAD_SIZE = 50 * 1024 * 1024  # 50 MB upload limit
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".pptx"}


def ensure_c_executable():
    """Ensure C executable exists and has execute permissions. Compiles if needed on POSIX/Linux."""
    if os.path.isfile(C_EXECUTABLE):
        if not sys.platform.startswith("win"):
            try:
                st = os.stat(C_EXECUTABLE)
                os.chmod(C_EXECUTABLE, st.st_mode | 0o755)
            except Exception as e:
                print(f"Warning: Could not set executable permissions on {C_EXECUTABLE}: {e}")
        return True

    src_dir = os.path.join(BASE_DIR, "src")
    include_dir = os.path.join(BASE_DIR, "include")
    if os.path.isdir(src_dir) and os.path.isdir(include_dir):
        print(f"C engine executable '{C_EXECUTABLE}' not found. Compiling from C source files...")
        import glob
        c_files = glob.glob(os.path.join(src_dir, "*.c"))
        if c_files:
            cmd = [
                "gcc", "-O2", "-Wall", "-Wextra", "-Wpedantic", "-Wshadow", "-Wconversion",
                "-std=c11", f"-I{include_dir}", "-o", C_EXECUTABLE
            ] + c_files
            try:
                subprocess.run(cmd, capture_output=True, text=True, check=True)
                if not sys.platform.startswith("win"):
                    st = os.stat(C_EXECUTABLE)
                    os.chmod(C_EXECUTABLE, st.st_mode | 0o755)
                print(f"Successfully compiled C engine: {C_EXECUTABLE}")
                return True
            except Exception as err:
                print(f"Automatic C engine compilation failed: {err}")
    print(f"Warning: C executable '{C_EXECUTABLE}' not found. Compilation with gcc may be required.")
    return False


def extract_text_from_pdf(filepath):
    try:
        doc = pymupdf.open(filepath)
    except Exception as e:
        raise ValueError("The PDF file appears to be corrupted or invalid.") from e

    if doc.is_encrypted:
        doc.close()
        raise ValueError("This PDF is password-protected and cannot be analyzed without a password.")

    text_parts = []
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            if text and text.strip():
                text_parts.append(text.strip())
    finally:
        doc.close()

    full_text = "\n\n".join(text_parts).strip()
    if not full_text:
        raise ValueError("This PDF does not contain extractable text. Scanned/image-only PDFs are not supported.")
    return full_text


def extract_text_from_docx(filepath):
    try:
        doc = docx.Document(filepath)
    except Exception as e:
        raise ValueError("The DOCX file appears to be corrupted or invalid.") from e

    text_parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            text_parts.append(t)

    for table in doc.tables:
        for row in table.rows:
            row_items = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_items:
                text_parts.append(" ".join(row_items))

    full_text = "\n\n".join(text_parts).strip()
    if not full_text:
        raise ValueError("This document does not contain extractable text.")
    return full_text


def extract_text_from_pptx(filepath):
    try:
        prs = pptx.Presentation(filepath)
    except Exception as e:
        raise ValueError("The PPTX file appears to be corrupted or invalid.") from e

    text_parts = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        slide_lines.append(t)
        if slide_lines:
            text_parts.append("\n".join(slide_lines))

    full_text = "\n\n".join(text_parts).strip()
    if not full_text:
        raise ValueError("This document does not contain extractable text.")
    return full_text


def extract_text_from_txt(filepath):
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            text = f.read().strip()
    except UnicodeDecodeError:
        with open(filepath, "r", encoding="latin-1") as f:
            text = f.read().strip()

    if not text:
        raise ValueError("This document does not contain extractable text.")
    return text


def extract_document_text(filepath, ext):
    ext = ext.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext == ".docx":
        return extract_text_from_docx(filepath)
    elif ext == ".pptx":
        return extract_text_from_pptx(filepath)
    elif ext == ".txt":
        return extract_text_from_txt(filepath)
    else:
        raise ValueError("Unsupported file format. Supported formats: TXT, PDF, DOCX, PPTX.")


class PlagiarismRequestHandler(http.server.BaseHTTPRequestHandler):
    def end_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(204)
        self.end_headers()

    def send_json(self, data, status=200):
        body = json.dumps(data).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path

        if path in ("/api/health", "/health", "/healthz"):
            self.send_json({
                "status": "healthy",
                "engine": "ready" if os.path.isfile(C_EXECUTABLE) else "missing",
                "executable": os.path.basename(C_EXECUTABLE)
            })
            return

        if path == "/api/sample-docs":
            self.handle_sample_docs()
            return

        if path == "/" or path == "":
            path = "/index.html"

        rel_path = os.path.normpath(path.lstrip("/"))
        file_path = os.path.join(UI_DIR, rel_path)

        if os.path.commonpath([UI_DIR, file_path]) == UI_DIR and os.path.isfile(file_path):
            content_type, _ = mimetypes.guess_type(file_path)
            if not content_type:
                content_type = "application/octet-stream"

            with open(file_path, "rb") as f:
                content = f.read()

            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        else:
            self.send_response(404)
            self.send_header("Content-Type", "text/plain")
            self.end_headers()
            self.wfile.write(b"404 Not Found")

    def do_POST(self):
        parsed = urlparse(self.path)
        path = parsed.path

        content_length = int(self.headers.get("Content-Length", 0))
        if content_length > MAX_UPLOAD_SIZE:
            self.send_json({"success": False, "error": "File size exceeds the 50 MB limit."}, 413)
            return

        post_body = self.rfile.read(content_length)

        try:
            data = json.loads(post_body.decode("utf-8")) if post_body else {}
        except Exception as e:
            self.send_json({"success": False, "error": f"Invalid JSON body: {str(e)}"}, 400)
            return

        if path == "/api/upload":
            self.handle_upload(data)
        elif path == "/api/compare":
            self.handle_compare(data)
        elif path == "/api/batch-compare":
            self.handle_batch_compare(data)
        elif path == "/api/export-report":
            self.handle_export_report(data)
        else:
            self.send_json({"success": False, "error": "Endpoint not found"}, 404)

    def handle_sample_docs(self):
        samples = []
        if os.path.isdir(DOCS_DIR):
            for fname in sorted(os.listdir(DOCS_DIR)):
                ext = os.path.splitext(fname)[1].lower()
                if ext in SUPPORTED_EXTENSIONS:
                    fpath = os.path.join(DOCS_DIR, fname)
                    samples.append({
                        "filename": fname,
                        "filepath": fpath.replace("\\", "/"),
                        "size": os.path.getsize(fpath)
                    })
        self.send_json({"success": True, "samples": samples})

    def handle_upload(self, data):
        filename = data.get("filename")
        content = data.get("content", "")
        is_base64 = data.get("is_base64", False)

        if not filename:
            self.send_json({"success": False, "error": "Missing filename"}, 400)
            return

        safe_name = os.path.basename(filename)
        ext = os.path.splitext(safe_name)[1].lower()

        if ext not in SUPPORTED_EXTENSIONS:
            self.send_json({
                "success": False,
                "error": "Unsupported file format. Supported formats: TXT, PDF, DOCX, PPTX."
            }, 400)
            return

        # Decode content bytes
        try:
            if is_base64:
                # Strip potential data URL prefix (e.g. data:application/pdf;base64,...)
                if "," in content:
                    content = content.split(",", 1)[1]
                content_bytes = base64.b64decode(content)
            else:
                content_bytes = content.encode("utf-8")
        except Exception as e:
            self.send_json({"success": False, "error": f"Failed to decode uploaded data: {str(e)}"}, 400)
            return

        if len(content_bytes) == 0:
            self.send_json({"success": False, "error": f"Document '{safe_name}' is empty."}, 400)
            return

        # Generate unique collision-safe identifier based on content hash
        file_hash = hashlib.sha256(content_bytes).hexdigest()[:12]
        base_name_no_ext = os.path.splitext(safe_name)[0]
        unique_original_name = f"{file_hash}_{safe_name}"
        unique_extracted_name = f"{file_hash}_{base_name_no_ext}.txt"

        original_path = os.path.join(ORIGINALS_DIR, unique_original_name)
        extracted_path = os.path.join(EXTRACTED_DIR, unique_extracted_name)

        # 1. Save original file to disk
        try:
            with open(original_path, "wb") as f:
                f.write(content_bytes)
        except Exception as e:
            self.send_json({"success": False, "error": f"Could not save original file: {str(e)}"}, 500)
            return

        # 2. Extract plain text according to format
        try:
            plain_text = extract_document_text(original_path, ext)
        except ValueError as ve:
            self.send_json({"success": False, "error": str(ve)}, 400)
            return
        except Exception as e:
            self.send_json({"success": False, "error": f"Extraction failure: {str(e)}"}, 500)
            return

        # 3. Save normalized plain text
        try:
            with open(extracted_path, "w", encoding="utf-8") as f:
                f.write(plain_text)
        except Exception as e:
            self.send_json({"success": False, "error": f"Could not save extracted text: {str(e)}"}, 500)
            return

        # 4. Query C engine for metadata
        cmd = [C_EXECUTABLE, "--doc-info", extracted_path]
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", check=True)
            stdout_text = (res.stdout or "").strip()
            doc_info = json.loads(stdout_text)
            doc_info["filepath"] = extracted_path.replace("\\", "/")
            doc_info["original_path"] = original_path.replace("\\", "/")
            doc_info["filename"] = safe_name  # Preserve original filename
            doc_info["file_size"] = len(content_bytes)
            doc_info["format"] = ext.lstrip(".").upper()
            self.send_json(doc_info)
        except Exception as e:
            self.send_json({
                "success": True,
                "filepath": extracted_path.replace("\\", "/"),
                "original_path": original_path.replace("\\", "/") if 'original_path' in locals() else "",
                "filename": safe_name,
                "file_size": len(content_bytes),
                "word_count": len(plain_text.split()),
                "sentence_count": plain_text.count(".") + plain_text.count("?") + plain_text.count("!"),
                "unique_words": len(set(plain_text.lower().split())),
                "top_words": [],
                "format": ext.lstrip(".").upper()
            })

    def handle_compare(self, data):
        ref_file = data.get("ref_file")
        cand_file = data.get("cand_file")
        ref_display_name = data.get("ref_name")
        cand_display_name = data.get("cand_name")

        if not ref_file or not cand_file:
            self.send_json({"success": False, "error": "Both ref_file and cand_file are required."}, 400)
            return

        cmd = [C_EXECUTABLE, "--compare-json", ref_file, cand_file]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
            if proc.returncode != 0:
                err_msg = (proc.stderr or proc.stdout or f"C engine process failed with return code {proc.returncode}").strip()
                self.send_json({"success": False, "error": err_msg}, 500)
                return

            stdout_text = (proc.stdout or "").strip()
            if not stdout_text:
                self.send_json({"success": False, "error": "C engine returned an empty response."}, 500)
                return

            result = json.loads(stdout_text)

            # Restore original user-facing filenames
            if ref_display_name:
                result["reference_name"] = ref_display_name
            if cand_display_name:
                result["compared_name"] = cand_display_name

            self.send_json(result)
        except subprocess.TimeoutExpired:
            self.send_json({"success": False, "error": "C engine comparison timed out."}, 500)
        except Exception as e:
            self.send_json({"success": False, "error": f"Failed to execute C engine: {str(e)}"}, 500)

    def handle_batch_compare(self, data):
        ref_file = data.get("ref_file")
        cand_files = data.get("cand_files", [])
        ref_display_name = data.get("ref_name")
        cand_display_map = data.get("cand_names", {})

        if not ref_file or not cand_files:
            self.send_json({"success": False, "error": "ref_file and at least one cand_file are required."}, 400)
            return

        cmd = [C_EXECUTABLE, "--batch-json", ref_file] + cand_files
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
            if proc.returncode != 0:
                err_msg = (proc.stderr or proc.stdout or f"C batch engine failed with return code {proc.returncode}").strip()
                self.send_json({"success": False, "error": err_msg}, 500)
                return

            stdout_text = (proc.stdout or "").strip()
            if not stdout_text:
                self.send_json({"success": False, "error": "C engine returned an empty response."}, 500)
                return

            result = json.loads(stdout_text)

            if ref_display_name:
                result["reference_name"] = ref_display_name

            # Map extracted filenames back to original user filenames
            for row in result.get("rankings", []):
                extracted_path = row.get("filepath", "")
                if extracted_path in cand_display_map:
                    row["document"] = cand_display_map[extracted_path]

            self.send_json(result)
        except Exception as e:
            self.send_json({"success": False, "error": f"Failed to execute batch comparison: {str(e)}"}, 500)

    def handle_export_report(self, data):
        ref_file = data.get("ref_file")
        cand_file = data.get("cand_file")
        out_name = data.get("output_filename")
        ref_display_name = data.get("ref_name")
        cand_display_name = data.get("cand_name")

        if not ref_file or not cand_file or not out_name:
            self.send_json({"success": False, "error": "Missing required fields for report export"}, 400)
            return

        safe_out_name = os.path.basename(out_name)
        if not safe_out_name.endswith(".txt"):
            safe_out_name += ".txt"

        dest_report = os.path.join(REPORTS_DIR, safe_out_name)
        cmd = [C_EXECUTABLE, "--export-report", ref_file, cand_file, dest_report]

        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=15)
            if proc.returncode != 0:
                err_msg = (proc.stderr or proc.stdout or f"C report engine failed with return code {proc.returncode}").strip()
                self.send_json({"success": False, "error": err_msg}, 500)
                return

            report_text = ""
            if os.path.isfile(dest_report):
                with open(dest_report, "r", encoding="utf-8") as f:
                    report_text = f.read()

                # Clean header display names in the saved report file
                if ref_display_name or cand_display_name:
                    lines = report_text.splitlines()
                    new_lines = []
                    for line in lines:
                        if line.startswith("Reference Document:") and ref_display_name:
                            new_lines.append("Reference Document:\n" + ref_display_name)
                        elif line.startswith("Compared Document:") and cand_display_name:
                            new_lines.append("Compared Document:\n" + cand_display_name)
                        else:
                            new_lines.append(line)
                    report_text = "\n".join(new_lines)
                    with open(dest_report, "w", encoding="utf-8") as f:
                        f.write(report_text)

            self.send_json({
                "success": True,
                "exported_path": dest_report.replace("\\", "/"),
                "filename": safe_out_name,
                "content": report_text
            })
        except Exception as e:
            self.send_json({"success": False, "error": f"Report export failed: {str(e)}"}, 500)


class ThreadedHTTPServer(socketserver.ThreadingMixIn, http.server.HTTPServer):
    daemon_threads = True
    allow_reuse_address = True


def get_port():
    """Resolve port from CLI argument, PORT environment variable (Render), or default."""
    # 1. Command-line argument: python server.py <port>
    if len(sys.argv) > 1:
        try:
            return int(sys.argv[1])
        except ValueError:
            pass
    # 2. Render / Cloud environment variable
    env_port = os.environ.get("PORT")
    if env_port:
        try:
            return int(env_port)
        except ValueError:
            pass
    # 3. Fallback for local development
    return DEFAULT_PORT


def run_server(host=HOST, port=None):
    if port is None:
        port = get_port()
    ensure_c_executable()
    server_address = (host, port)
    with ThreadedHTTPServer(server_address, PlagiarismRequestHandler) as httpd:
        print(f"Plagiarism Detection Server running on http://{host}:{port}")
        print(f"Loaded C Engine executable: {C_EXECUTABLE}")
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nServer shutting down gracefully.")
            httpd.shutdown()


if __name__ == "__main__":
    run_server()
